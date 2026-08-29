from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


NAVY = "17365D"
BLUE = "D9EAF7"


def create_trade_report(journal, path: str, account=None) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    trades, snapshots = journal.trades(), journal.snapshots()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False
    summary.append(["IG Demo Trading Report"])
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary.merge_cells("A1:D1")
    summary.append(["Generated (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)])
    summary.append(["Recorded trade events", len(trades)])
    summary.append(["Current balance", float(account.balance) if account else None])
    summary.append(["Current equity", float(account.equity) if account else None])
    summary.append(["Open profit/loss", float(account.profit) if account else None])
    for row in range(2, 7):
        summary.cell(row, 1).font = Font(bold=True)
    summary["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    for cell in ("B4", "B5", "B6"):
        summary[cell].number_format = '#,##0.00;[Red](#,##0.00);-'
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 22

    trade_sheet = wb.create_sheet("Trade Events")
    headers = ["Time (UTC)", "Symbol", "Side", "Volume", "Price", "Stop Loss", "Take Profit", "Probability", "Ticket", "Status", "Message"]
    trade_sheet.append(headers)
    for item in trades:
        trade_sheet.append([item["created_at"], item["symbol"], item["side"], item["volume"], item["price"],
                            item["stop_loss"], item["take_profit"], item["probability"], item["ticket"], item["status"], item["message"]])
    _format_table(trade_sheet, len(headers))
    trade_sheet.freeze_panes = "A2"
    trade_sheet.auto_filter.ref = trade_sheet.dimensions
    trade_sheet.column_dimensions["A"].width = 27
    trade_sheet.column_dimensions["K"].width = 42

    balance_sheet = wb.create_sheet("Balance History")
    balance_sheet.append(["Time (UTC)", "Balance", "Equity", "Margin", "Free Margin"])
    for item in snapshots:
        balance_sheet.append([item["created_at"], item["balance"], item["equity"], item["margin"], item["free_margin"]])
    _format_table(balance_sheet, 5)
    balance_sheet.freeze_panes = "A2"
    balance_sheet.column_dimensions["A"].width = 27
    if snapshots:
        chart = LineChart()
        chart.title = "Balance and Equity History"
        chart.y_axis.title = "Account currency"
        chart.x_axis.title = "Snapshot"
        chart.add_data(Reference(balance_sheet, min_col=2, max_col=3, min_row=1, max_row=len(snapshots)+1), titles_from_data=True)
        chart.set_categories(Reference(balance_sheet, min_col=1, min_row=2, max_row=len(snapshots)+1))
        chart.height, chart.width = 8, 16
        summary.add_chart(chart, "A9")
    wb.save(output)
    return output


def _format_table(sheet, columns: int) -> None:
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")
    for col in range(2, columns + 1):
        sheet.column_dimensions[chr(64 + col)].width = 15
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in {4, 5, 6, 7}:
                cell.number_format = '#,##0.00000;[Red](#,##0.00000);-'
            elif cell.column == 8:
                cell.number_format = "0.0%"
